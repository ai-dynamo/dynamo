# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0

"""Final evaluation tables (every method) from published Match Config results.

For each results JSON (one per topology) the report is reduced, on the single
evaluation metric ``goodput_per_gpu`` (good requests per second divided by the
average GPU count, i.e. good requests per GPU-second), to:

* a summary table with one row per method and every column reported on two
  workload sets as ``all (golden)``: mean goodput per GPU-second, mean rank and
  pairwise win rate (from ``rank_tables``), mean good rate and mean average
  GPUs. ``all`` is every workload in the matrix; ``golden`` is the last six
  workloads in matrix order (the Astra GPT-OSS Golden Set);
* (workbook only) the wide tables: per-workload ``goodput_per_gpu`` (sheet
  ``<mode> goodput_per_gpu``), and the rank matrix on that same metric (sheet
  ``<mode> ranks``) with one row per method, one column per workload and each
  cell the method's rank on that workload (1 = best, ties share the average
  rank), both with mean columns over all and over golden workloads, plus the
  long-form per-method x per-workload data behind every number.

Metric values are averaged over repetitions first, then over the workloads of
each set with equal weight. When several results are given, a cross-mode table
of the CloudAI RL entries (``cloudai-rl-*``) is appended to the markdown, with
the best non-RL method of each mode for reference.

The markdown (console and ``--out``) carries the summary tables; ``--xlsx``
writes the workbook (``openpyxl``, static values computed here; regenerate it
rather than editing cells). With ``--out`` and no ``--xlsx`` the workbook is
written next to the markdown with an ``.xlsx`` suffix.

    python scripts/evaluation_table.py \\
        --results agg=runs/match-sim-agg-all-datasets/results.agg.json \\
        --results disagg=runs/match-sim-disagg-all-datasets/results.disagg.json \\
        --out runs/evaluation.md            # also writes runs/evaluation.xlsx
"""

from __future__ import annotations

import argparse
import json
import statistics
import sys
from collections.abc import Mapping, Sequence
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from autoscaling_arena.rank_tables import (
    GOLDEN_SET_SIZE,
    WORKLOAD_SETS,
    average_ranks,
    methods_for_metric,
    metric_table,
    rank_table,
    select_workloads,
    sla_order,
    workload_order,
)

METRIC = "goodput_per_gpu"
RL_PREFIX = "cloudai-rl"
# Extra per-run metrics copied into the long-form data sheets (the leaderboard
# HTML columns), with their sheet labels and number formats. goodput_rps is
# left out on purpose: goodput_per_gpu is the metric.
EXTRA_METRICS: dict[str, tuple[str, str]] = {
    "mean_ttft_ms": ("Mean TTFT (ms)", "0"),
    "p99_ttft_ms": ("P99 TTFT (ms)", "0"),
    "mean_itl_ms": ("Mean ITL / TPOT (ms)", "0.0"),
    "oscillation_count": ("Oscillations", "0.0"),
    "scale_events": ("Scale events", "0.0"),
}
# (column title, aggregate key, markdown format, Excel number format)
COLUMNS: tuple[tuple[str, str, str, str], ...] = (
    ("Goodput/GPU-s", "goodput_per_gpu", "{:.3f}", "0.000"),
    ("Mean rank", "mean_rank", "{:.2f}", "0.00"),
    ("Win rate", "win_rate", "{:.3f}", "0.000"),
    ("Good rate", "good_rate", "{:.1%}", "0.0%"),
    ("Avg GPUs", "avg_gpus", "{:.2f}", "0.00"),
)
HOW_TO_READ = [
    "## How to read these tables",
    "",
    "- **One row per autoscaler, one table per topology.** `agg` runs a single pool of aggregated "
    "workers; `disagg` runs separate prefill and decode pools. The seven methods, the 14 traces, the "
    "SLO and the simulator are identical within a table, so rows are directly comparable.",
    "- **The metric is `goodput_per_gpu`**: SLO-compliant (good) requests served per second, divided by "
    "the average number of GPUs the autoscaler had provisioned (starting, active and draining workers "
    "all count). It is good requests per GPU-second, i.e. how much useful work one GPU-second bought; "
    "higher is better. Multiply by 3600 for good requests per GPU-hour.",
    "- **`all (golden)`**: every cell shows the mean over all 14 traces, then in parentheses the mean "
    "over the six Astra GPT-OSS Golden Set traces (hour-long production-like traffic). The other "
    "eight are shorter registry traces (the Mooncake anchor plus seven synthetic patterns). Each "
    "trace is first averaged over its 3 repetitions, then traces are averaged with equal weight.",
    "- **Goodput/GPU-s** is the headline column and orders the rows. **Mean rank** ranks the methods on "
    "each trace by that metric (1 = best, ties share the average rank) and averages the ranks, so it "
    "rewards being consistently good rather than winning a few traces by a wide margin. **Win rate** "
    "is the share of (trace, opponent) head-to-head comparisons a method wins (a tie counts half); "
    "0.5 means average.",
    "- **Good rate** (share of all requests that met both the TTFT and ITL SLO) and **Avg GPUs** "
    "explain *how* a method got its goodput per GPU: a high good rate on a large fleet (static) and a "
    "lower good rate on a small fleet (lean autoscalers) can score alike. Neither is a ranking "
    "criterion.",
    "- **Fixed fleets are the reference points.** The static entry (4 workers, or 4 prefill + 4 decode) "
    "has the highest good rate because it never under-provisions, and the lowest goodput per GPU "
    "because it never scales down.",
    "- **Caveat for the RL row.** Five of the six golden traces were training workloads for the RL "
    "planner (only random-bursts, flash_crowd and diurnal were held out), so its golden numbers are "
    "partly in-sample; the other methods never see the traces in advance.",
    "",
]


def _mean(values: Sequence[float]) -> float:
    return statistics.fmean(values) if values else float("nan")


def _methods(report: Mapping[str, Any]) -> list[str]:
    return methods_for_metric(
        list(
            dict.fromkeys(
                str(r["autoscaler"]) for r in report.get("matrix") or report["results"]
            )
        ),
        METRIC,
    )


def _samples(
    report: Mapping[str, Any], *, sla: str
) -> dict[tuple[str, str], dict[str, list[float]]]:
    """Per (method, workload) lists of per-repetition values over ok rows of one SLA."""
    methods = set(_methods(report))
    workloads = set(workload_order(report))
    out: dict[tuple[str, str], dict[str, list[float]]] = {}
    for row in report["results"]:
        if row.get("status") != "ok" or row.get("sla") != sla:
            continue
        key = (str(row.get("autoscaler")), str(row.get("workload")))
        if key[0] not in methods or key[1] not in workloads:
            continue
        metrics = row.get("metrics") or {}
        raw = row.get("raw_metrics") or {}
        entry = out.setdefault(
            key, {"goodput_per_gpu": [], "good_rate": [], "avg_gpus": []}
        )
        entry["goodput_per_gpu"].append(float(metrics[METRIC]))
        entry["good_rate"].append(float(metrics.get("good_rate", float("nan"))))
        entry["avg_gpus"].append(float(raw.get("average_gpus", float("nan"))))
        for name in _extra_metrics(report):
            value = metrics.get(name)
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                entry.setdefault(name, []).append(float(value))
    return out


def _extra_metrics(report: Mapping[str, Any]) -> list[str]:
    """The leaderboard's extra metric columns (``summary.metrics``) in EXTRA_METRICS order."""
    listed = {str(m) for m in ((report.get("summary") or {}).get("metrics") or [])}
    return [name for name in EXTRA_METRICS if name in listed]


def method_rows(report: Mapping[str, Any], *, sla: str) -> dict[str, dict[str, Any]]:
    """Per-method aggregates on each workload set, over ok rows of one SLA target.

    ``rows[method]["per_workload"][workload]`` is the repetition-averaged
    goodput per GPU-second; ``rows[method][set][key]`` the set-level aggregate
    for ``set`` in ``("all", "golden")`` and ``key`` one of the COLUMNS keys.
    """
    samples = _samples(report, sla=sla)
    ranks = {
        ws: {
            r["method"]: r
            for r in rank_table(report, metric=METRIC, workload_set=ws, sla=sla)["rows"]
        }
        for ws in WORKLOAD_SETS
    }
    out: dict[str, dict[str, Any]] = {}
    for m in _methods(report):
        per_workload = {
            w: {f: _mean(v[f]) for f in ("goodput_per_gpu", "good_rate", "avg_gpus")}
            for (mm, w), v in samples.items()
            if mm == m and v["goodput_per_gpu"]
        }
        entry: dict[str, Any] = {
            "per_workload": {w: v["goodput_per_gpu"] for w, v in per_workload.items()}
        }
        for ws in WORKLOAD_SETS:
            chosen = [w for w in select_workloads(report, ws) if w in per_workload]
            entry[ws] = {
                "goodput_per_gpu": _mean(
                    [per_workload[w]["goodput_per_gpu"] for w in chosen]
                ),
                "good_rate": _mean([per_workload[w]["good_rate"] for w in chosen]),
                "avg_gpus": _mean([per_workload[w]["avg_gpus"] for w in chosen]),
                "mean_rank": ranks[ws].get(m, {}).get("mean_rank", float("nan")),
                "win_rate": ranks[ws].get(m, {}).get("pairwise_win_rate", float("nan")),
            }
        out[m] = entry
    return out


def rank_matrix(report: Mapping[str, Any], *, sla: str) -> dict[str, dict[str, float]]:
    """``{method: {workload: rank}}`` by repetition-averaged goodput per GPU-second.

    Rank 1 = best per workload; tied values share the average rank. Only
    workloads where every method has an ok repetition are ranked.
    """
    workloads = workload_order(report)
    methods = _methods(report)
    table = metric_table(
        report["results"], metric=METRIC, sla=sla, workloads=workloads, methods=methods
    )
    matrix: dict[str, dict[str, float]] = {m: {} for m in methods}
    for workload in workloads:
        if workload not in table:
            continue
        for method, rank in average_ranks(table[workload]).items():
            matrix[method][workload] = rank
    return matrix


def _cell(entry: Mapping[str, Any], key: str, fmt: str) -> str:
    return f"{fmt.format(entry['all'][key])} ({fmt.format(entry['golden'][key])})"


def format_mode(label: str, report: Mapping[str, Any], *, sla: str) -> list[str]:
    rows = method_rows(report, sla=sla)
    workloads = workload_order(report)
    order = sorted(rows, key=lambda m: -rows[m]["all"]["goodput_per_gpu"])
    reps = len({r.get("repetition") for r in report["results"]})
    lines = [
        f"## {label} mode",
        "",
        f"{len(workloads)} workloads x {reps} repetitions, SLA `{sla}`. Every column reads "
        "`all (golden)`: the mean over all workloads, then over the last "
        f"{GOLDEN_SET_SIZE} workloads in matrix order (the Astra GPT-OSS Golden Set). "
        "Goodput/GPU-s is `goodput_per_gpu`: good requests per second divided by the average GPU "
        "count, i.e. good requests per GPU-second (x3600 for GPU-hours). Values are averaged over "
        "repetitions, then over workloads with equal weight. Mean rank 1 = best per workload (ties "
        "share the average rank); pairwise win rate counts ties as half.",
        "",
        "| Method | " + " | ".join(title for title, _, _, _ in COLUMNS) + " |",
        "|---|" + "---:|" * len(COLUMNS),
    ]
    for m in order:
        name = f"**{m}**" if m.startswith(RL_PREFIX) else m
        lines.append(
            f"| {name} | "
            + " | ".join(_cell(rows[m], key, fmt) for _, key, fmt, _ in COLUMNS)
            + " |"
        )
    lines.append("")
    return lines


def format_rl_summary(
    modes: Mapping[str, Mapping[str, Any]], *, slas: Mapping[str, str]
) -> list[str]:
    columns = list(
        dict.fromkeys(
            m for r in modes.values() for m in _methods(r) if m.startswith(RL_PREFIX)
        )
    )
    lines = [
        "## CloudAI RL planner by mode",
        "",
        "Mean goodput per GPU-second, `all (golden)`; the last column is the best non-RL method of the mode.",
        "",
        "| Mode | " + " | ".join(columns) + " | Best other method |",
        "|---|" + "---:|" * len(columns) + "---|",
    ]
    for label, report in modes.items():
        rows = method_rows(report, sla=slas[label])
        cells = [
            "n/a" if c not in rows else _cell(rows[c], "goodput_per_gpu", "{:.3f}")
            for c in columns
        ]
        others = [m for m in rows if not m.startswith(RL_PREFIX)]
        best = (
            max(others, key=lambda m: rows[m]["all"]["goodput_per_gpu"])
            if others
            else None
        )
        best_cell = (
            "n/a"
            if best is None
            else f"{best}: {_cell(rows[best], 'goodput_per_gpu', '{:.3f}')}"
        )
        lines.append(f"| {label} | " + " | ".join(cells) + f" | {best_cell} |")
    lines.append("")
    return lines


def _load(
    results: Mapping[str, Path], *, sla: str | None
) -> tuple[dict[str, Mapping[str, Any]], dict[str, str]]:
    modes: dict[str, Mapping[str, Any]] = {}
    slas: dict[str, str] = {}
    for label, path in results.items():
        report = json.loads(Path(path).read_text())
        modes[label] = report
        slas[label] = sla or sla_order(report)[0]
    return modes, slas


def render(
    results: Mapping[str, Path], *, sla: str | None = None, workbook: Path | None = None
) -> str:
    modes, slas = _load(results, sla=sla)
    lines = ["# Autoscaler evaluation: goodput per GPU-second", "", *HOW_TO_READ]
    if workbook is not None:
        lines += [
            f"Per-workload goodput_per_gpu, the rank matrices and the long-form data (with the "
            f"leaderboard's latency and scaling metrics) are in `{workbook.name}` "
            "(sheets `<mode> goodput_per_gpu`, `<mode> ranks`, `<mode> data`).",
            "",
        ]
    for label, report in modes.items():
        lines += format_mode(label, report, sla=slas[label])
    has_rl = any(
        m.startswith(RL_PREFIX) for label, r in modes.items() for m in _methods(r)
    )
    if len(modes) > 1 or has_rl:
        lines += format_rl_summary(modes, slas=slas)
    return "\n".join(lines)


def write_workbook(
    results: Mapping[str, Path], path: Path, *, sla: str | None = None
) -> Path:
    """Write the full tables to an ``.xlsx`` workbook (static values, Arial)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    modes, slas = _load(results, sla=sla)
    base = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    title_font = Font(name="Arial", size=12, bold=True)
    header_fill = PatternFill("solid", fgColor="DDEBF7")
    golden_fill = PatternFill("solid", fgColor="FFF2CC")
    right = Alignment(horizontal="right")

    def header(ws, values: Sequence[str], *, golden: set[str] = frozenset()) -> None:
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=1, column=col, value=value)
            cell.font = bold
            cell.fill = golden_fill if value in golden else header_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=True)

    def widths(ws, first: int, others: int, count: int) -> None:
        ws.column_dimensions["A"].width = first
        for col in range(2, count + 1):
            ws.column_dimensions[get_column_letter(col)].width = others

    def put(
        ws,
        row: int,
        col: int,
        value: Any,
        *,
        fmt: str | None = None,
        strong: bool = False,
    ) -> None:
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = bold if strong else base
        cell.alignment = right
        if fmt:
            cell.number_format = fmt

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    notes: list[tuple[str, Any]] = [
        ("Autoscaler evaluation: goodput per GPU-second", title_font),
        (
            f"Generated {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')} by "
            "autoscaling_arena.evaluation_table (scripts/evaluation_table.py). Static values; regenerate "
            "the file instead of editing cells.",
            base,
        ),
        ("", base),
        (
            "Metric: goodput_per_gpu = good requests per second / average GPU count = good requests per "
            "GPU-second (x3600 for GPU-hours). A request is good when it meets every SLO threshold of the "
            "SLA target.",
            base,
        ),
        (
            "Repetitions are averaged first, then workloads with equal weight. 'golden' = the last six "
            "workloads in matrix order (Astra GPT-OSS Golden Set, yellow headers); 'all' = every workload.",
            base,
        ),
        (
            "Rank 1 = best per workload; tied values share the average rank. Pairwise win rate = share of "
            "(workload, opponent) comparisons won, ties count 0.5 (both from autoscaling_arena.rank_tables).",
            base,
        ),
        ("", base),
        (
            "Sheets per mode: '<mode> summary' (one row per method, all and golden columns), "
            "'<mode> goodput_per_gpu' (method x workload goodput_per_gpu = good requests per GPU-second, best per "
            "workload in bold), '<mode> ranks' (method x workload rank BY goodput_per_gpu, rank 1 in bold), "
            "'<mode> data' (long form: method, workload, set, goodput_per_gpu, good rate, average GPUs, "
            "the leaderboard's latency and scaling metrics, repetitions; every value is the mean over "
            "repetitions).",
            base,
        ),
        ("", base),
    ]
    for label, report in modes.items():
        provenance = report.get("provenance") or {}
        notes.append(
            (
                f"{label}: {results[label]} | SLA {slas[label]} | session {provenance.get('session_id', '')} | "
                f"finished {provenance.get('finished_at', '')}",
                base,
            )
        )
    for row, (text, font) in enumerate(notes, start=1):
        cell = readme.cell(row=row, column=1, value=text)
        cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    readme.column_dimensions["A"].width = 140

    for label, report in modes.items():
        s = slas[label]
        rows = method_rows(report, sla=s)
        workloads = workload_order(report)
        golden = set(select_workloads(report, "golden"))
        order = sorted(rows, key=lambda m: -rows[m]["all"]["goodput_per_gpu"])

        # --- summary ---
        ws = wb.create_sheet(f"{label} summary")
        titles = ["Method"]
        for title, _, _, _ in COLUMNS:
            titles += [f"{title} (all)", f"{title} (golden)"]
        header(ws, titles)
        for r, m in enumerate(order, start=2):
            ws.cell(row=r, column=1, value=m).font = (
                bold if m.startswith(RL_PREFIX) else base
            )
            col = 2
            for _, key, _, fmt in COLUMNS:
                for ws_name in WORKLOAD_SETS:
                    put(ws, r, col, rows[m][ws_name][key], fmt=fmt)
                    col += 1
        ws.freeze_panes = "B2"
        widths(ws, 22, 13, len(titles))

        # --- goodput matrix ---
        ws = wb.create_sheet(f"{label} goodput_per_gpu")
        header(
            ws,
            ["Method (goodput_per_gpu)", *workloads, "Mean (all)", "Mean (golden)"],
            golden=golden,
        )
        best = {
            w: max(rows[m]["per_workload"].get(w, float("-inf")) for m in order)
            for w in workloads
        }
        for r, m in enumerate(order, start=2):
            ws.cell(row=r, column=1, value=m).font = (
                bold if m.startswith(RL_PREFIX) else base
            )
            for c, w in enumerate(workloads, start=2):
                value = rows[m]["per_workload"].get(w)
                put(
                    ws,
                    r,
                    c,
                    value,
                    fmt="0.000",
                    strong=value is not None and value == best[w],
                )
            put(
                ws,
                r,
                len(workloads) + 2,
                rows[m]["all"]["goodput_per_gpu"],
                fmt="0.000",
                strong=True,
            )
            put(
                ws,
                r,
                len(workloads) + 3,
                rows[m]["golden"]["goodput_per_gpu"],
                fmt="0.000",
                strong=True,
            )
        ws.freeze_panes = "B2"
        widths(ws, 30, 11, len(workloads) + 3)

        # --- rank matrix ---
        matrix = rank_matrix(report, sla=s)
        ranked = [w for w in workloads if any(w in ranks for ranks in matrix.values())]
        ws = wb.create_sheet(f"{label} ranks")
        header(
            ws,
            [
                "Method (rank by goodput_per_gpu)",
                *ranked,
                "Mean (all)",
                "Mean (golden)",
            ],
            golden=golden,
        )
        rank_order = sorted(
            matrix,
            key=lambda m: (
                statistics.fmean(matrix[m].values()) if matrix[m] else float("inf"),
                m,
            ),
        )
        for r, m in enumerate(rank_order, start=2):
            ws.cell(row=r, column=1, value=m).font = (
                bold if m.startswith(RL_PREFIX) else base
            )
            for c, w in enumerate(ranked, start=2):
                value = matrix[m].get(w)
                put(ws, r, c, value, fmt="0.#", strong=value == 1)
            values_all = [matrix[m][w] for w in ranked if w in matrix[m]]
            values_golden = [
                matrix[m][w] for w in ranked if w in matrix[m] and w in golden
            ]
            put(
                ws,
                r,
                len(ranked) + 2,
                _mean(values_all) if values_all else None,
                fmt="0.00",
                strong=True,
            )
            put(
                ws,
                r,
                len(ranked) + 3,
                _mean(values_golden) if values_golden else None,
                fmt="0.00",
                strong=True,
            )
        ws.freeze_panes = "B2"
        widths(ws, 30, 11, len(ranked) + 3)

        # --- long-form data ---
        ws = wb.create_sheet(f"{label} data")
        extras = _extra_metrics(report)
        header(
            ws,
            [
                "Method",
                "Workload",
                "Set",
                "goodput_per_gpu (good req / GPU-s)",
                "Good rate",
                "Avg GPUs",
                *(EXTRA_METRICS[name][0] for name in extras),
                "Repetitions",
            ],
        )
        samples = _samples(report, sla=s)
        r = 2
        for m in order:
            for w in workloads:
                entry = samples.get((m, w))
                if entry is None:
                    continue
                ws.cell(row=r, column=1, value=m).font = base
                ws.cell(row=r, column=2, value=w).font = base
                ws.cell(
                    row=r, column=3, value="golden" if w in golden else "registry"
                ).font = base
                put(ws, r, 4, _mean(entry["goodput_per_gpu"]), fmt="0.000")
                put(ws, r, 5, _mean(entry["good_rate"]), fmt="0.0%")
                put(ws, r, 6, _mean(entry["avg_gpus"]), fmt="0.00")
                col = 7
                for name in extras:
                    values = entry.get(name, [])
                    put(
                        ws,
                        r,
                        col,
                        _mean(values) if values else None,
                        fmt=EXTRA_METRICS[name][1],
                    )
                    col += 1
                put(ws, r, col, len(entry["goodput_per_gpu"]))
                r += 1
        ws.freeze_panes = "A2"
        last = get_column_letter(7 + len(extras))
        ws.auto_filter.ref = f"A1:{last}{max(r - 1, 1)}"
        for col, width in zip("ABCDEF", (22, 24, 10, 18, 11, 10)):
            ws.column_dimensions[col].width = width
        for index in range(7, 8 + len(extras)):
            ws.column_dimensions[get_column_letter(index)].width = 14

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "--results",
        action="append",
        required=True,
        metavar="LABEL=PATH",
        help="published results JSON with a mode label, e.g. agg=runs/.../results.agg.json (repeatable)",
    )
    parser.add_argument(
        "--sla", help="SLA target name (default: the first one in each matrix)"
    )
    parser.add_argument("--out", help="also write the markdown to this file")
    parser.add_argument(
        "--xlsx",
        help="write the full tables (per-workload goodput, rank matrices, data) to this workbook; "
        "defaults to --out with an .xlsx suffix",
    )
    parser.add_argument(
        "--no-xlsx",
        action="store_true",
        help="skip the workbook even when --out is given",
    )
    args = parser.parse_args(argv)
    results: dict[str, Path] = {}
    for item in args.results:
        label, _, path = item.partition("=")
        if not path:
            parser.error(f"--results expects LABEL=PATH, got {item!r}")
        results[label] = Path(path)
    workbook: Path | None = None
    if not args.no_xlsx:
        if args.xlsx:
            workbook = Path(args.xlsx)
        elif args.out:
            workbook = Path(args.out).with_suffix(".xlsx")
    text = render(results, sla=args.sla, workbook=workbook)
    if args.out:
        Path(args.out).write_text(text + "\n")
    if workbook is not None:
        write_workbook(results, workbook, sla=args.sla)
    sys.stdout.write(text + "\n")
    if workbook is not None:
        sys.stdout.write(f"\nWorkbook: {workbook}\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
